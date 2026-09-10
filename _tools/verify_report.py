# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r"C:\Users\Lenovo\Desktop\Web前端设计与开发实训\大作业\大作业报告-青阅书城-Web前端设计与开发实训.docx")
print("总段落数:", len(doc.paragraphs))
print("总表格数:", len(doc.tables))
for i in [8, 10, 11, 12]:
    print("封面段%d: %s" % (i, doc.paragraphs[i].text[:60]))
full_text = " ".join([p.text for p in doc.paragraphs])
keywords = ["明暗主题", "限时秒杀", "优惠券", "图书对比", "快速预览",
            "搜索自动补全", "物流时间线", "编辑推荐", "回到顶部",
            "3D倾斜", "懒加载", "软工2507", "8002125223", "邹浏源"]
print("\n关键词检查:")
for kw in keywords:
    status = "OK" if kw in full_text else "MISSING!"
    print("  %s: %s" % (kw, status))
if doc.tables:
    print("\n测试用例表行数:", len(doc.tables[0].rows), "(含表头)")
print("\n验证完成!")
